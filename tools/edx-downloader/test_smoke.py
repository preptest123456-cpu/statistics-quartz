#!/usr/bin/env python3
"""Offline smoke tests for edx_course_downloader.

Exercises everything that does not need a browser or an edX session:
filename sanitising, URL parsing, the block tree, sjson->srt conversion,
HTML->text, and a real Word document build (including tall-screenshot slicing).

    python test_smoke.py
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import edx_course_downloader as ed  # noqa: E402

FAILURES: list[str] = []


def check(label: str, condition: bool, detail: str = "") -> None:
    status = "PASS" if condition else "FAIL"
    print(f"  [{status}] {label}" + (f"  -- {detail}" if detail and not condition else ""))
    if not condition:
        FAILURES.append(label)


def test_sanitize() -> None:
    print("\nsanitize_component")
    check("colon becomes ' -'",
          ed.sanitize_component("edex_China-West Relations: Dilemmas and Lessons")
          == "edex_China-West Relations - Dilemmas and Lessons")
    check("illegal characters stripped",
          ed.sanitize_component('a<b>c:d"e|f?g*h') == "abc -defgh",
          ed.sanitize_component('a<b>c:d"e|f?g*h'))
    check("slashes flattened", "/" not in ed.sanitize_component("week 1/2"))
    check("reserved names escaped", ed.sanitize_component("CON").startswith("_"))
    check("empty falls back", ed.sanitize_component("   ") == "untitled")
    check("length capped", len(ed.sanitize_component("x" * 500)) <= 90)


def test_url_parsing() -> None:
    print("\ncourse_id_from_url")
    url = (
        "https://learning.edx.org/course/course-v1:KingsCollegeLondon+SSPP_STCx4+1T2025/"
        "block-v1:KingsCollegeLondon+SSPP_STCx4+1T2025+type@sequential+block@886a432c/"
        "block-v1:KingsCollegeLondon+SSPP_STCx4+1T2025+type@vertical+block@ef55329f"
    )
    check("course id extracted from a deep unit URL",
          ed.course_id_from_url(url) == "course-v1:KingsCollegeLondon+SSPP_STCx4+1T2025",
          str(ed.course_id_from_url(url)))
    check("no id in a non-course URL", ed.course_id_from_url("https://edx.org/dashboard") is None)


def sample_tree() -> ed.Block:
    blocks = {
        "course@root": {"type": "course", "display_name": "China-West Relations",
                        "children": ["chapter@1"]},
        "chapter@1": {"type": "chapter", "display_name": "Week 1: Foundations",
                      "children": ["seq@1"]},
        "seq@1": {"type": "sequential", "display_name": "Lecture block",
                  "children": ["vert@1"]},
        "vert@1": {"type": "vertical", "display_name": "Introduction",
                   "children": ["html@1", "video@1"]},
        "html@1": {"type": "html", "display_name": "Reading",
                   "student_view_data": {"html": "<p>Hello</p>"}},
        "video@1": {"type": "video", "display_name": "Welcome",
                    "student_view_data": {
                        "duration": 754.0,
                        "transcripts": {"en": "/transcript/en"},
                        "encoded_videos": {
                            "hls": {"url": "https://example.invalid/v.m3u8", "file_size": 0},
                            "mobile_low": {"url": "https://example.invalid/low.mp4",
                                           "file_size": 10_000_000},
                            "desktop_mp4": {"url": "https://example.invalid/hd.mp4",
                                            "file_size": 90_000_000},
                        },
                    }},
    }
    return ed.build_tree(blocks, "course@root")


def test_tree() -> None:
    print("\nbuild_tree")
    tree = sample_tree()
    check("root type", tree.block_type == "course")
    check("node count", sum(1 for _ in tree.walk()) == 6)
    verticals = tree.descendants_of_type(["vertical"])
    check("one vertical found", len(verticals) == 1)
    unit = verticals[0]
    check("breadcrumb", unit.breadcrumb() == "Week 1: Foundations / Lecture block / Introduction",
          unit.breadcrumb())
    check("folder name numbering", unit.folder_name() == "01 - Introduction", unit.folder_name())
    check("videos located", len(tree.descendants_of_type(ed.VIDEO_TYPES)) == 1)
    check("parents wired", unit.parent is not None and unit.parent.block_type == "sequential")


def test_video_selection() -> None:
    print("\nVideoDownloader.resolve")
    tree = sample_tree()
    video = tree.descendants_of_type(ed.VIDEO_TYPES)[0]

    cfg = ed.Config()
    downloader = ed.VideoDownloader.__new__(ed.VideoDownloader)   # no network/browser needed
    downloader.cfg = cfg
    downloader.s = None
    downloader.manifest = None
    downloader.limiter = ed.RateLimiter(0)
    downloader._ffmpeg = None
    downloader._ytdlp = None

    best = downloader.resolve(video, Path("/tmp"))
    check("best profile picked", best is not None and best.profile == "desktop_mp4",
          best.profile if best else "None")
    check("mp4 extension", best.path.suffix == ".mp4")
    check("duration carried", best.duration == 754.0)
    check("transcripts carried", best.transcripts == {"en": "/transcript/en"})

    cfg.quality = "mobile_low"
    picked = downloader.resolve(video, Path("/tmp"))
    check("explicit profile honoured", picked.profile == "mobile_low", picked.profile)

    empty = ed.Block(block_id="v2", block_type="video", display_name="No sources")
    check("no sources -> None", downloader.resolve(empty, Path("/tmp")) is None)


def test_sjson() -> None:
    print("\nsjson_to_srt")
    srt = ed.sjson_to_srt('{"start":[0,2500],"end":[2000,4000],"text":["First line","Second line"]}')
    check("two cues", srt.count("-->") == 2)
    check("timestamp format", "00:00:00,000 --> 00:00:02,000" in srt, srt.splitlines()[1])
    check("second cue offset", "00:00:02,500 --> 00:00:04,000" in srt)
    check("non-json passes through", ed.sjson_to_srt("1\n00:00 --> 00:01\nhi") .startswith("1"))


def test_html_to_text() -> None:
    print("\nhtml_to_text")
    text = ed.html_to_text("<div><h2>Title</h2><p>One<br>Two</p><ul><li>A</li><li>B</li></ul></div>")
    for token in ("Title", "One", "Two", "A", "B"):
        check(f"contains {token!r}", token in text)
    check("tags removed", "<" not in text)


def test_docx_build() -> None:
    print("\nDocxBuilder")
    from PIL import Image

    tree = sample_tree()
    cfg = ed.Config()
    with tempfile.TemporaryDirectory() as tmp:
        tmpdir = Path(tmp)
        cfg.root = tmpdir

        short = tmpdir / "short.png"
        Image.new("RGB", (1280, 700), (240, 244, 250)).save(short)
        tall = tmpdir / "tall.png"
        Image.new("RGB", (1280, 5200), (250, 250, 250)).save(tall)

        builder = ed.DocxBuilder(cfg, tree)
        chapter = tree.children[0]
        sequential = chapter.children[0]
        unit = sequential.children[0]
        builder.add_section(chapter)
        builder.add_subsection(sequential)

        capture = ed.UnitCapture(block=unit, directory=tmpdir)
        capture.html = (
            "<div><h2>Key ideas</h2>"
            "<p>Some <strong>bold</strong> and <em>italic</em> text with a "
            "<a href='https://example.invalid/paper'>link</a>.</p>"
            "<ul><li>First point</li><li>Second point</li></ul>"
            "<table><tr><th>Year</th><th>Event</th></tr>"
            "<tr><td>1972</td><td>Visit</td></tr></table>"
            "<blockquote>A quotation.</blockquote></div>"
        )
        capture.screenshot = short
        asset = ed.VideoAsset(block=unit, url="https://example.invalid/hd.mp4",
                              profile="desktop_mp4", path=tmpdir / "Welcome.mp4",
                              size=90_000_000, duration=754.0)
        asset.downloaded = True
        asset.transcript_paths = {"en": tmpdir / "Welcome.en.srt"}
        capture.videos = [asset]
        builder.add_unit(capture)

        tall_capture = ed.UnitCapture(block=unit, directory=tmpdir)
        tall_capture.text = "Plain text unit.\n\nSecond paragraph."
        tall_capture.screenshot = tall
        builder.add_unit(tall_capture)

        out = builder.save(tmpdir / "course.docx")
        check("document written", out.is_file() and out.stat().st_size > 10_000,
              f"{out.stat().st_size if out.is_file() else 0} bytes")

        import docx
        doc = docx.Document(str(out))
        body = "\n".join(p.text for p in doc.paragraphs)
        check("course title present", "China-West Relations" in body)
        check("section heading present", "Week 1: Foundations" in body)
        check("unit text rendered", "Second point" in body)
        check("video listed", "Welcome.mp4" in body)
        check("breadcrumb present", "Week 1: Foundations / Lecture block" in body)
        check("table rendered", len(doc.tables) >= 1)
        check("images embedded",
              len([r for r in doc.part.rels.values() if "image" in r.reltype]) >= 2,
              str(len([r for r in doc.part.rels.values() if "image" in r.reltype])))


def test_manifest() -> None:
    print("\nManifest")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "manifest.json"
        manifest = ed.Manifest(path)
        check("unknown key not done", not manifest.done("b1", "unit"))
        manifest.record("b1", "unit", title="Intro")
        manifest.fail("b2", "video", "boom")
        manifest.save()
        check("file written", path.is_file())
        reloaded = ed.Manifest(path)
        check("state survives a reload", reloaded.done("b1", "unit"))
        check("failures are not 'done'", not reloaded.done("b2", "video"))


def test_misc() -> None:
    print("\nhelpers")
    check("human_size", ed.human_size(1536) == "1.5 KB", ed.human_size(1536))
    check("human_duration h:mm:ss", ed.human_duration(3725) == "1:02:05", ed.human_duration(3725))
    check("human_duration m:ss", ed.human_duration(75) == "1:15")
    check("rate limiter is a no-op at 0", ed.RateLimiter(0).wait() is None)

    cfg = ed.parse_args(["--no-videos", "--limit", "3", "--url",
                         "https://learning.edx.org/course/course-v1:Org+Num+1T2025/x/y"])
    check("CLI parses course id from URL", cfg.course_id == "course-v1:Org+Num+1T2025")
    check("CLI --no-videos", cfg.want_videos is False)
    check("CLI --limit", cfg.limit == 3)
    check("CLI default folder name", cfg.folder_name == ed.DEFAULT_FOLDER_NAME)


def test_compiled_dependency_preflight() -> None:
    print("\ncompiled dependency preflight")
    imported: list[str] = []

    def working_importer(module: str) -> object:
        imported.append(module)
        return object()

    ed.preflight_compiled_dependencies(working_importer)
    check("imports every compiled dependency",
          imported == ["lxml.etree", "charset_normalizer"], str(imported))

    attempted: list[str] = []

    def sac_blocked_importer(module: str) -> object:
        attempted.append(module)
        if module == "lxml.etree":
            raise ImportError(
                "DLL load failed while importing etree: "
                "An Application Control policy has blocked this file."
            )
        if module == "charset_normalizer":
            raise ImportError(
                "DLL load failed while importing md: "
                "An Application Control policy has blocked this file."
            )
        return object()

    try:
        ed.preflight_compiled_dependencies(sac_blocked_importer)
    except ed.ArchiveError as exc:
        message = str(exc)
        check("checks all dependencies after one failure",
              attempted == ["lxml.etree", "charset_normalizer"], str(attempted))
        check("explains Smart App Control", "Smart App Control" in message)
        check("does not recommend disabling SAC", "Do not disable" in message)
        check("shows the system-site-packages workaround",
              "--system-site-packages" in message)
        check("reports lxml failure", "lxml.etree" in message)
        check("reports charset-normalizer failure", "charset_normalizer" in message)
    else:
        check("blocked imports fail before archiving", False, "ArchiveError was not raised")

    original_preflight = ed.preflight_compiled_dependencies
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp) / "must-not-be-created"

        def fail_preflight() -> None:
            raise ed.ArchiveError("simulated SAC block")

        try:
            ed.preflight_compiled_dependencies = fail_preflight
            result = ed.main([
                "--out", str(Path(tmp)),
                "--folder-name", root.name,
                "--dry-run",
            ])
        finally:
            ed.preflight_compiled_dependencies = original_preflight
        check("main exits on a failed preflight", result == 2, str(result))
        check("preflight runs before output creation", not root.exists(), str(root))


def main() -> int:
    print("edX course downloader -- offline smoke tests")
    for test in (test_sanitize, test_url_parsing, test_tree, test_video_selection,
                 test_sjson, test_html_to_text, test_docx_build, test_manifest, test_misc,
                 test_compiled_dependency_preflight):
        test()
    print("\n" + "-" * 60)
    if FAILURES:
        print(f"{len(FAILURES)} check(s) FAILED:")
        for name in FAILURES:
            print(f"  - {name}")
        return 1
    print("All checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
